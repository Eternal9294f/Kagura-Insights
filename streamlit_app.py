import os
import json
import streamlit as st
import pandas as pd
from datetime import datetime
import io

from src.analysis import validate_and_load_reviews, filter_by_weeks, analyze_corpus
from src.generator import generate_weekly_pulse, generate_fee_explainer
from src.mcp_client import KaguraMCPClient

# Attempt imports for PDF and DOCX export
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

def create_pdf(text: str, title: str) -> bytes:
    if not HAS_REPORTLAB:
        return b""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=54, leftMargin=54, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]
    normal_style.fontSize = 10
    normal_style.leading = 14
    
    story = []
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        spaceAfter=12
    )
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 10))

    for line in text.split("\n"):
        clean = line.strip()
        if not clean:
            story.append(Spacer(1, 6))
            continue
        clean_escaped = clean.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(clean_escaped, normal_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    return buffer.getvalue()

def create_docx(text: str, title: str) -> bytes:
    if not HAS_DOCX:
        return b""
    doc = docx.Document()
    doc.add_heading(title, 0)
    for line in text.split("\n"):
        clean = line.strip()
        if clean:
            doc.add_paragraph(clean)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

st.set_page_config(
    page_title="Project Kagura — Groww Review Intelligence",
    page_icon="🔥",
    layout="wide"
)

st.title("🔥 Project Kagura: Review Intelligence & MCP HITL Workflow")
st.caption("NextLeap PM Fellowship Milestone 4 | Product Intelligence & Customer Clarification Engine for Groww")

with st.sidebar:
    st.header("⚙️ System Configuration")
    product_name = st.text_input("Target Product", value="Groww")
    weeks_window = st.slider("Analysis Window (Weeks)", min_value=8, max_value=12, value=10)
    
    st.divider()
    st.markdown("### 📥 Dataset Selection")
    uploaded_file = st.file_uploader("Upload Public Reviews CSV", type=["csv"], help="Expected columns: review_id, date, rating, source, review_text")
    
    st.divider()
    st.markdown("### 🛡️ Security & MCP Guardrails")
    st.success("Human-in-the-Loop: ACTIVE\nDraft Mode: LOCAL PERSISTENCE\nAuto-send: DISABLED")
    
    st.divider()
    st.markdown("### 📚 Authoritative Corpus")
    st.markdown("- **Entity**: Groww India")
    st.markdown("- **Focus**: Mutual Fund Exit Loads")
    st.markdown("- **Benchmark**: 60 Verified Reviews (`data/reviews.csv`)")

tab_intel, tab_pulse, tab_explainer, tab_mcp, tab_sources = st.tabs([
    "📊 Review Intelligence", "📝 Weekly Pulse", "💡 Fee Explainer", "🔒 MCP Approval Gate", "📖 Sources & Ledger"
])

# Load either uploaded file or verified benchmark dataset
default_csv = os.path.join(os.path.dirname(__file__), "data", "reviews.csv")
active_csv = uploaded_file if uploaded_file is not None else default_csv

df_raw, status_msg = validate_and_load_reviews(active_csv)
df_filtered = filter_by_weeks(df_raw, weeks=weeks_window)
analysis = analyze_corpus(df_filtered)

pulse_text, pulse_words = generate_weekly_pulse(analysis, product_name=product_name)
explainer = generate_fee_explainer(analysis, product_name=product_name)

mcp_client = KaguraMCPClient()
staged = mcp_client.stage_payloads(pulse_text, explainer, analysis, product_name)

with tab_intel:
    st.subheader("1. Ingestion & Clustering Overview")
    st.info(f"📁 Dataset Status: {status_msg}")
    
    col1, col2, col3, col4 = st.columns(4)
    col1.metric("Total Reviews Processed", analysis["total_reviews"])
    col2.metric("Date Span", f"{analysis.get('date_min', 'N/A')} to {analysis.get('date_max', 'N/A')}")
    col3.metric("Canonical Themes", len(analysis["theme_counts"]))
    col4.metric("Fee Confusion Triggered", "YES" if analysis["fee_confusion_detected"] else "NO")

    st.divider()
    c_left, c_right = st.columns([1, 1])
    with c_left:
        st.markdown("#### Thematic Distribution (Strictly ≤ 5 Categories)")
        theme_rows = [{"Theme": k, "Count": v, "Share": f"{analysis['theme_percentages'].get(k, 0)}%"} for k, v in analysis["theme_counts"].items()]
        st.dataframe(pd.DataFrame(theme_rows), use_container_width=True, hide_index=True)

    with c_right:
        st.markdown("#### Genuine Verbatim Quotes (Direct Slices)")
        for i, q in enumerate(analysis["verbatim_quotes"], 1):
            st.info(f"**Quote {i}:** \"{q}\"")

    st.markdown("#### Isolated Fee Confusion")
    if analysis["fee_confusion_detected"]:
        st.warning(analysis["fee_confusion_detail"])
    else:
        st.info("No recurring fee confusion identified in the loaded corpus.")

with tab_pulse:
    st.subheader(f"2. Weekly Product Pulse ({pulse_words} words / Max 250 words)")
    st.markdown(pulse_text)
    if pulse_words <= 250:
        st.success("✅ Milestone Compliance: Pulse is strictly under the 250-word cap.")
    else:
        st.error(f"⚠️ Non-compliant: Exceeds 250 words ({pulse_words} words).")
    
    st.divider()
    st.markdown("#### Export Report")
    d_c1, d_c2, d_c3 = st.columns(3)
    d_c1.download_button("⬇️ Download Markdown", pulse_text, "weekly-pulse.md", "text/markdown", use_container_width=True)
    if HAS_REPORTLAB:
        d_c2.download_button("⬇️ Download PDF", create_pdf(pulse_text, f"Weekly Product Pulse — {product_name}"), "weekly-pulse.pdf", "application/pdf", use_container_width=True)
    if HAS_DOCX:
        d_c3.download_button("⬇️ Download DOCX", create_docx(pulse_text, f"Weekly Product Pulse — {product_name}"), "weekly-pulse.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

with tab_explainer:
    st.subheader("3. Customer Clarification (Fee Explainer)")
    st.markdown(f"### {explainer['title']}")
    for bullet in explainer["bullets"]:
        st.markdown(f"- {bullet}")
    st.markdown(f"**Last checked:** `{explainer['last_checked']}`")
    st.markdown("**Authoritative Source References:**")
    for src in explainer["source_links"]:
        st.markdown(f"- [{src['title']}]({src['url']})")
    
    st.divider()
    explainer_full_md = (
        f"# {explainer['title']}\n\nLast checked: {explainer['last_checked']}\n\n"
        + "\n".join([f"- {b}" for b in explainer["bullets"]])
        + "\n\n### Official Sources:\n"
        + "\n".join([f"- {s['title']}: {s['url']}" for s in explainer["source_links"]])
    )
    e_c1, e_c2, e_c3 = st.columns(3)
    e_c1.download_button("⬇️ Download Markdown", explainer_full_md, "fee-explainer.md", "text/markdown", use_container_width=True)
    if HAS_REPORTLAB:
        e_c2.download_button("⬇️ Download PDF", create_pdf(explainer_full_md, explainer['title']), "fee-explainer.pdf", "application/pdf", use_container_width=True)
    if HAS_DOCX:
        e_c3.download_button("⬇️ Download DOCX", create_docx(explainer_full_md, explainer['title']), "fee-explainer.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)

with tab_mcp:
    st.subheader("4. Approval-Gated MCP Tool Workflow")
    st.markdown("Inspect staged tool execution requests awaiting explicit operator confirmation:")

    with st.expander("🔍 View Staged Tool Payloads", expanded=True):
        st.markdown("**Tool 1: `append_note` Staged Payload (JSON)**")
        st.json(staged["notes_payload"])
        st.markdown("**Tool 2: `create_draft` Staged Payload (Email Draft)**")
        st.code(f"Subject: {staged['subject']}\nTo: product-team@groww.in\n\n{staged['email_body']}", language="markdown")

    if "approval_state" not in st.session_state:
        st.session_state["approval_state"] = "PENDING"

    b_col1, b_col2, _ = st.columns([1, 1, 2])
    with b_col1:
        if st.button("✅ Approve & Dispatch MCP Tools", type="primary", use_container_width=True):
            res_n, res_d = mcp_client.execute_approved_workflow(staged)
            st.session_state["approval_state"] = "APPROVED"
            st.session_state["mcp_results"] = (res_n, res_d)

    with b_col2:
        if st.button("❌ Deny / Halt Execution", use_container_width=True):
            st.session_state["approval_state"] = "DENIED"

    st.divider()
    if st.session_state["approval_state"] == "APPROVED":
        st.success("🎉 MCP Action Authorized: Payloads dispatched and persisted to local records!")
        rn, rd = st.session_state["mcp_results"]
        st.write(f"- Appended to Notes: `{rn['destination']}`")
        st.write(f"- Staged Email Draft: `{rd['destinations'][0]}` and `{rd['destinations'][1]}`")
        
        # Download buttons for executed MCP artifacts
        st.markdown("#### Download MCP Artifacts")
        m_c1, m_c2 = st.columns(2)
        with open(rn['destination'], 'r', encoding='utf-8') as f:
            note_content = f.read()
        with open(rd['destinations'][0], 'r', encoding='utf-8') as f:
            email_content = f.read()
        m_c1.download_button("⬇️ Download doc-entry.json", note_content, "doc-entry.json", "application/json", use_container_width=True)
        m_c2.download_button("⬇️ Download email-draft.md", email_content, "email-draft.md", "text/markdown", use_container_width=True)

    elif st.session_state["approval_state"] == "DENIED":
        st.error("🛑 Action Denied: MCP workflow aborted. No downstream records were modified.")
    else:
        st.info("⏳ Awaiting human authorization. No MCP mutations have occurred.")

with tab_sources:
    st.subheader("5. Corpus Metadata & Authoritative Ledger")
    sources_df = pd.read_csv(os.path.join(os.path.dirname(__file__), "sources.csv"))
    st.dataframe(sources_df, use_container_width=True, hide_index=True)
    with open(os.path.join(os.path.dirname(__file__), "data", "corpus-meta.json"), "r") as f:
        meta_json = json.load(f)
    st.json(meta_json)
